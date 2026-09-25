#!/usr/bin/env python3
"""Recalculate a workbook and report what an analyst checks before handing it over.

    python3 check_workbook.py model.xlsx

Needs openpyxl and formulas (pip install openpyxl formulas). The workbook is
recalculated in Python, so this works with no spreadsheet app installed.

Reports:
  1. error values after recalculation (#DIV/0!, #REF!, #VALUE!, #NAME?, #N/A, #NUM!)
  2. the Checks sheet, row by row: non-zero flags are failures, except under a
     heading that says "Warnings", which are reported but don't fail
  3. the Cover sheet, so the headline numbers can be read back
  4. lint: defined names Excel reads as cell addresses, rows whose formula
     changes across periods, numbers typed into formulas, and TRUE/FALSE flags
     written as --(...)

Exits 1 if any cell holds an error value, any error check is non-zero, or a
defined name is really a cell address.
"""
import contextlib
import io
import re
import sys
from collections import Counter, defaultdict

try:
    import formulas
    import openpyxl
    from openpyxl.formula.tokenizer import Token, Tokenizer
    from openpyxl.formula.translate import Translator
    from openpyxl.utils import column_index_from_string
except ImportError:
    sys.exit("needs openpyxl and formulas: pip install openpyxl formulas")

ERRORS = {"#DIV/0!", "#REF!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
# Unit and structural constants that may appear inside formulas.
ALLOWED_NUMBERS = {0, 1, 2, 4, 7, 12, 52, 100, 360, 365, 1000}
LIMIT = 15  # findings shown per section
RESERVED_COLUMNS = 4  # label, units, total/constant, spacer: periods start after these


def recalculate(path):
    """Return {(sheet, 'A1'): value} for every calculated cell."""
    names = {n.upper(): n for n in openpyxl.load_workbook(path, read_only=True).sheetnames}
    with contextlib.redirect_stderr(io.StringIO()), contextlib.redirect_stdout(io.StringIO()):
        solution = formulas.ExcelModel().loads(path).finish().calculate()
    values = {}
    for key, rng in solution.items():
        m = re.match(r"^'?\[[^\]]+\](.+?)'?!\$?([A-Z]+)\$?(\d+)$", key)
        value = getattr(rng, "value", None)
        if not m or getattr(value, "shape", None) != (1, 1):
            continue
        sheet = names.get(m.group(1).upper(), m.group(1))
        v = value[0, 0]
        v = v.item() if hasattr(v, "item") else v
        values[(sheet, f"{m.group(2)}{m.group(3)}")] = v
    return values


def is_error(v):
    return str(v) in ERRORS


def show(v):
    if isinstance(v, bool):
        return str(v).upper()
    if isinstance(v, float):
        return f"{v:,.4f}" if abs(v) < 10 else f"{v:,.2f}"
    return str(v)


def sheet_rows(ws, values):
    """Yield (row, label, [computed values]) for rows that have a text label."""
    for row in ws.iter_rows():
        label, vals = None, []
        for cell in row:
            v = values.get((ws.title, cell.coordinate), cell.value)
            if v is None or v == "":
                continue
            if label is None and isinstance(cell.value, str) and not cell.value.startswith("="):
                label = cell.value.strip()
            elif label is not None:
                vals.append(v)
        if label:
            yield row[0].row, label, vals


def cell_like_names(wb):
    """Names such as NOI2 or CoC1 are cell addresses to Excel, so =NOI2 reads an empty cell."""
    bad = []
    for name in wb.defined_names:
        m = re.fullmatch(r"([A-Za-z]{1,3})(\d+)", name)
        if m and column_index_from_string(m.group(1).upper()) <= 16384 and int(m.group(2)) <= 1048576:
            bad.append(f"{name}: Excel reads it as cell {m.group(1).upper()}{m.group(2)}; rename it (e.g. {m.group(1)}_Y{m.group(2)})")
        elif re.fullmatch(r"[Rr]\d*[Cc]\d*", name):
            bad.append(f"{name}: Excel reads it as an R1C1 reference; rename it")
    return bad


def canonical(formula, origin, target):
    try:
        return Translator(formula, origin=origin).translate_formula(target)
    except Exception:
        return formula


def lint(wb):
    broken_rows, typed_numbers, bool_flags = [], [], []
    for ws in wb.worksheets:
        by_row = defaultdict(list)
        for row in ws.iter_rows():
            for cell in row:
                f = cell.value
                if not (isinstance(f, str) and f.startswith("=")):
                    continue
                by_row[cell.row].append(cell)
                if re.search(r"--\s*\(|--\s*AND\(|--\s*OR\(", f, re.I):
                    bool_flags.append(f"{ws.title}!{cell.coordinate}: {f}")
                try:
                    tokens = Tokenizer(f).items
                except Exception:
                    continue
                for t in tokens:
                    if t.type == Token.OPERAND and t.subtype == Token.NUMBER:
                        try:
                            n = float(t.value)
                        except ValueError:
                            continue
                        tolerance = ws.title.lower() == "checks" and abs(n) <= 1
                        if n not in ALLOWED_NUMBERS and not tolerance:
                            typed_numbers.append(f"{ws.title}!{cell.coordinate}: {t.value} in {f}")
        # A time-series row carries one formula across every period: translate each
        # cell's formula to the row's first formula column and compare.
        for r, cells in by_row.items():
            runs, run = [], [cells[0]]
            for prev, cell in zip(cells, cells[1:]):
                if cell.column == prev.column + 1:
                    run.append(cell)
                else:
                    runs.append(run)
                    run = [cell]
            runs.append(run)
            cells = [c for c in max(runs, key=len) if c.column > RESERVED_COLUMNS]
            if len(cells) < 4:
                continue
            first = cells[0].coordinate
            forms = {c.coordinate: canonical(c.value, c.coordinate, first) for c in cells}
            common, count = Counter(forms.values()).most_common(1)[0]
            if count >= 3 and count >= 0.6 * len(cells) and count < len(cells):
                odd = [c for c, f in forms.items() if f != common]
                broken_rows.append(f"{ws.title}!row {r}: {', '.join(odd[:6])}{' …' if len(odd) > 6 else ''} differ from the row's formula")
    return broken_rows, typed_numbers, bool_flags


def section(title, items, empty):
    print(f"\n## {title}")
    if not items:
        print(f"  {empty}")
    for line in items[:LIMIT]:
        print(f"  {line}")
    if len(items) > LIMIT:
        print(f"  … and {len(items) - LIMIT} more")


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    wb = openpyxl.load_workbook(path)
    values = recalculate(path)
    print(f"# {path}: {len(wb.sheetnames)} sheets, {len(values)} cells recalculated")

    errors = [f"{s}!{a}: {v}" for (s, a), v in sorted(values.items()) if is_error(v)]
    section("Error values", errors, "none")

    failing, warnings = [], []
    checks = next((ws for ws in wb.worksheets if ws.title.lower() == "checks"), None)
    print("\n## Checks")
    if checks is None:
        print("  no sheet named Checks: every model needs one (see references/core/conventions.md)")
    else:
        in_warnings = False
        for r, label, vals in sheet_rows(checks, values):
            if "warning" in label.lower():
                in_warnings = in_warnings or not vals  # a heading row switches blocks
            nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
            bad = any(is_error(v) for v in vals) or any(abs(n) > 1e-9 for n in nums[:1])
            warn = in_warnings or "warning" in label.lower()
            if bad:
                (warnings if warn else failing).append(label)
            mark = ("!" if warn else "✗") if bad else " "
            print(f"  {mark} {label}: {', '.join(show(v) for v in vals[:4])}")
        print(f"  {len(failing)} failing checks" + (f": {'; '.join(failing[:LIMIT])}" if failing else ""))
        print(f"  {len(warnings)} warnings" + (f": {'; '.join(warnings[:LIMIT])}" if warnings else ""))

    cover = next((ws for ws in wb.worksheets if ws.title.lower() == "cover"), None)
    if cover is not None:
        print("\n## Cover")
        for r, label, vals in sheet_rows(cover, values):
            print(f"  {label}: {', '.join(show(v) for v in vals[:3])}")

    bad_names = cell_like_names(wb)
    section("Defined names that are really cell addresses", bad_names, "none")
    broken_rows, typed_numbers, bool_flags = lint(wb)
    section("Rows whose formula changes across periods", broken_rows, "none")
    section("Numbers typed into formulas (move each to an input)", typed_numbers, "none")
    section("TRUE/FALSE flags: write 1*(...) or IF(...,1,0), since --(...) sums to 0 in some engines",
            bool_flags, "none")

    ok = not errors and not failing and not bad_names
    print(f"\n{'PASS' if ok else 'FAIL'}: {len(errors)} error values, {len(failing)} failing checks, {len(bad_names)} cell-like names, {len(warnings)} warnings")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
