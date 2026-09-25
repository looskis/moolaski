#!/usr/bin/env python3
"""Regression test for skills/moolaski-financial-analyst/scripts/check_workbook.py.

Builds one clean workbook and one with every problem the script must catch,
then checks the verdicts. Needs openpyxl and formulas.

Usage: python3 scripts/test_check_workbook.py
"""
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

SCRIPT = Path(__file__).resolve().parent.parent / "skills/moolaski-financial-analyst/scripts/check_workbook.py"


def build(path, broken):
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    inputs = wb.create_sheet("Inputs")
    calc = wb.create_sheet("Calc")
    checks = wb.create_sheet("Checks")

    inputs["A1"], inputs["B1"] = "Rent", 1000
    inputs["A2"], inputs["B2"] = "Start", 1
    wb.defined_names["NOI2" if broken else "NOI_Y2"] = DefinedName(
        "NOI2" if broken else "NOI_Y2", attr_text="Calc!$F$3")

    # label, units, total, spacer, then periods 0..5 in E:J
    calc["A1"] = "Index"
    calc["A2"] = "Operating flag"
    calc["A3"] = "Rent"
    for i, col in enumerate("EFGHIJ"):
        calc[f"{col}1"] = i
        flag = f"--({col}1>=Inputs!$B$2)" if broken else f"1*({col}1>=Inputs!$B$2)"
        calc[f"{col}2"] = f"={flag}"
        rent = "=Inputs!$B$1*1.03" if broken and col == "H" else f"=Inputs!$B$1*{col}2"
        calc[f"{col}3"] = rent
    calc["C2"] = "=SUM(E2:J2)"
    calc["C3"] = "=SUM(E3:J3)"
    calc["C4"] = "=C3/(C2-C2)" if broken else "=C3/C2"

    checks["A1"] = "Errors"
    checks["A2"], checks["B2"] = "Flag sum matches its length", "=IFERROR(1*(ABS(Calc!C2-5)>0.01),1)"
    checks["A3"] = "Warnings"
    checks["A4"], checks["B4"] = "Rent below 2000", "=1*(Inputs!B1<2000)"

    cover["A1"], cover["B1"] = "Total rent", f"={'NOI2' if broken else 'NOI_Y2'}"
    wb.save(path)


def run(path):
    p = subprocess.run([sys.executable, str(SCRIPT), str(path)], capture_output=True, text=True)
    return p.returncode, p.stdout


def main():
    failures = []
    with tempfile.TemporaryDirectory() as d:
        clean, broken = Path(d, "clean.xlsx"), Path(d, "broken.xlsx")
        build(clean, broken=False)
        build(broken, broken=True)

        code, out = run(clean)
        if code != 0 or "PASS" not in out or "1 warnings" not in out:
            failures.append(f"clean workbook should pass with one warning:\n{out}")

        code, out = run(broken)
        expected = {
            "fails": code == 1,
            "error value": "Calc!C4: #DIV/0!" in out,
            "failing check (--() flags sum to 0)": "✗ Flag sum matches its length" in out,
            "cell-like name": "NOI2: Excel reads it as cell NOI2" in out,
            "row whose formula changes": "Calc!row 3: H3" in out,
            "typed number": "1.03 in" in out,
            "--() flag": "Calc!E2: =--(" in out,
        }
        failures += [f"broken workbook: missed {k}" for k, ok in expected.items() if not ok]
        if failures:
            failures.append(out)

    for f in failures:
        print(f"FAIL: {f}")
    print("ok" if not failures else f"{len(failures)} failures")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
